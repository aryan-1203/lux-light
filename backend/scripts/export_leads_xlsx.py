"""Export the Leads sheet to a formatted .xlsx file.

Usage (from backend/):
    python scripts/export_leads_xlsx.py --out leads.xlsx [--tab Leads] [--from-csv path.csv]

Without --from-csv it reads the sheet with the same credentials as the API
(GOOGLE_SERVICE_ACCOUNT_B64 + SHEET_ID from the environment or backend/.env).

Every cell uses Abadi 9 pt; the header row is bold, frozen and has an autofilter.
Note: Abadi is a Microsoft font (bundled with Microsoft 365). Excel/Numbers fall back to a
default font on machines where it is not installed; the file itself is unaffected.
"""

import argparse
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

FONT_NAME = "Abadi"
FONT_SIZE = 9
MIN_WIDTH, MAX_WIDTH = 8, 60


def read_csv(path: str) -> list[list[str]]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        return [row for row in csv.reader(f)]


def read_sheet(tab: str) -> list[list[str]]:
    from app.config import Settings
    from app.services.sheets_client import GspreadSheetsClient

    settings = Settings()
    if not settings.sheets_configured:
        raise SystemExit(
            "GOOGLE_SERVICE_ACCOUNT_B64 and LEADS_SHEET_ID / NEWSLETTER_SHEET_ID (or SHEET_ID) must be set"
            " (or use --from-csv)."
        )
    client = GspreadSheetsClient(settings.google_service_account_b64, sheet_ids=settings.spreadsheet_ids)
    return client.get_all_values(tab)


def write_xlsx(rows: list[list[str]], out: str, title: str = "Leads") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    body_font = Font(name=FONT_NAME, size=FONT_SIZE)
    header_font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)
    widths: dict[int, int] = {}

    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"  # keep text; never let a cell become a live formula
            cell.font = header_font if r == 1 else body_font
            widths[c] = max(widths.get(c, 0), len(str(value)))

    for c, width in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = max(MIN_WIDTH, min(MAX_WIDTH, width + 2))
    if rows:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Export the Leads sheet to .xlsx (Abadi 9 pt).")
    parser.add_argument("--out", required=True, help="output .xlsx path")
    parser.add_argument("--tab", default="Leads", help="worksheet tab to export (default: Leads)")
    parser.add_argument("--from-csv", dest="from_csv", help="read rows from a CSV file instead of Sheets")
    args = parser.parse_args(argv)

    rows = read_csv(args.from_csv) if args.from_csv else read_sheet(args.tab)
    write_xlsx(rows, args.out, title=args.tab)
    print(f"Wrote {max(len(rows) - 1, 0)} rows to {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
