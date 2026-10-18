import csv
import importlib.util
from pathlib import Path

from openpyxl import load_workbook

from app.services.rows import LEADS_HEADER

SCRIPT = Path(__file__).resolve().parent.parent / "scripts" / "export_leads_xlsx.py"


def _load_script():
    spec = importlib.util.spec_from_file_location("export_leads_xlsx", SCRIPT)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_export_from_csv(tmp_path) -> None:
    src = tmp_path / "leads.csv"
    rows = [
        LEADS_HEADER,
        ["2026-10-18 15:42:07", "LX-7F3K9Q", "Asha Rao", "asha@gmail.com", "+919876543210", "Bengaluru"]
        + [""] * 14
        + ["New"],
        ["2026-10-18 16:00:00", "LX-ABCDEF", "=cmd", "", "+918000000000", "Pune", "x" * 200] + [""] * 14,
    ]
    with src.open("w", newline="") as f:
        csv.writer(f).writerows(rows)
    out = tmp_path / "out" / "leads.xlsx"

    assert _load_script().main(["--out", str(out), "--from-csv", str(src)]) == 0

    ws = load_workbook(out).active
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == ws.dimensions
    assert [c.value for c in ws[1]] == LEADS_HEADER
    for row in ws.iter_rows():
        for cell in row:
            assert cell.font.name == "Abadi"
            assert cell.font.sz == 9
            assert bool(cell.font.b) is (cell.row == 1)
    assert ws["E2"].value == "+919876543210"
    assert ws["C3"].value == "=cmd" and ws["C3"].data_type == "s"
    widths = {k: v.width for k, v in ws.column_dimensions.items()}
    assert widths["G"] == 60  # clamped
    assert all(8 <= w <= 60 for w in widths.values())
